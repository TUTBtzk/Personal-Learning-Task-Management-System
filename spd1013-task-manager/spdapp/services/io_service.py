# -*- coding: utf-8 -*-
"""导入导出与数据备份。"""
import csv
import io
import os
import sqlite3
from datetime import datetime

from flask import current_app

from ..db import get_db
from ..models import course_repo, log_repo, task_repo
from .task_service import (PRIORITY_BY_LABEL, PRIORITY_LABELS, STATUS_BY_LABEL,
                           STATUS_LABELS, TITLE_MAX, is_date)

# 导出表头 = 导入模板表头，导出的文件可以直接改完再导回来
HEADERS = ["任务名称", "所属课程", "截止日期", "优先级", "完成状态", "备注",
           "创建时间", "完成时间"]


def _rows_to_matrix(rows):
    """数据库行 → 二维表格，枚举值转成中文，方便直接给人看。"""
    matrix = []
    for r in rows:
        matrix.append([
            r["title"],
            r["course_name"] or "未分类",
            r["due_date"],
            PRIORITY_LABELS[r["priority"]],
            STATUS_LABELS[r["status"]],
            r["note"] or "",
            r["created_at"] or "",
            r["completed_at"] or "",
        ])
    return matrix


def export_filename(ext):
    """导出文件名带项目标识与时间戳，截图时一眼能看出是本人的文件。"""
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    return f"{current_app.config['PROJECT_CODE']}_任务清单_{stamp}.{ext}"


def export_csv(rows):
    """导出 CSV。

    编码必须用 utf-8-sig（带 BOM）：少了 BOM，Excel 双击打开中文全是乱码。
    """
    buf = io.StringIO(newline="")
    writer = csv.writer(buf)
    writer.writerow(HEADERS)
    writer.writerows(_rows_to_matrix(rows))
    log_repo.add_log("export", "CSV", f"导出 {len(rows)} 条任务")
    return buf.getvalue().encode("utf-8-sig")


def export_xlsx(rows):
    """导出 Excel：表头加粗填色、冻结首行、列宽预设、开启筛选。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "任务清单"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0D6EFD")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in _rows_to_matrix(rows):
        ws.append(row)

    for idx, width in enumerate([34, 22, 13, 8, 10, 30, 20, 20], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    log_repo.add_log("export", "Excel", f"导出 {len(rows)} 条任务")
    return buf.getvalue()


def _read_csv(data):
    """CSV 编码依次按 utf-8-sig / utf-8 / gbk 尝试，兼容 Excel 另存的中文文件。"""
    for encoding in ("utf-8-sig", "utf-8", "gbk"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("文件编码无法识别，请另存为 UTF-8 或 GBK 编码的 CSV")
    return [row for row in csv.reader(io.StringIO(text))]


def _read_xlsx(data):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active
    return [["" if v is None else str(v).strip() for v in values]
            for values in ws.iter_rows(values_only=True)]


def import_tasks(filename, data):
    """导入任务：逐行校验，坏行跳过并汇报原因（部分成功，不整体回滚）。"""
    ext = os.path.splitext(filename or "")[1].lower()
    if ext not in current_app.config["ALLOWED_UPLOAD_EXT"]:
        raise ValueError("只支持 .csv 和 .xlsx 文件")

    rows = _read_csv(data) if ext == ".csv" else _read_xlsx(data)
    rows = [r for r in rows if any(str(c).strip() for c in r)]   # 丢掉纯空行
    if not rows:
        raise ValueError("文件里没有可导入的数据")

    start = 1 if str(rows[0][0]).strip() in ("任务名称", "title") else 0
    success, failed = 0, []
    for line_no, raw in enumerate(rows[start:], start=start + 1):
        item, reason = _parse_row(raw)
        if reason:
            failed.append({"row": line_no, "reason": reason,
                           "raw": " | ".join(str(c) for c in raw[:3])})
            continue
        task_repo.create_task(item)
        success += 1

    log_repo.add_log("import", os.path.basename(filename or ""),
                     f"成功 {success} 条，失败 {len(failed)} 条")
    return {"success": success, "failed": failed, "total": len(rows) - start}


def _parse_row(raw):
    """校验单行，返回 (任务数据, 错误原因)。两者必有一个为 None。"""
    cells = list(raw) + [""] * (6 - len(raw))

    title = str(cells[0] or "").strip()
    if not title:
        return None, "任务名称为空"
    if len(title) > TITLE_MAX:
        return None, f"任务名称超过 {TITLE_MAX} 个字"

    # Excel 里的日期读出来可能是 2026-09-09 00:00:00，截前 10 位即可
    due = str(cells[2] or "").strip()[:10]
    if not is_date(due):
        return None, f"截止日期「{cells[2]}」不是 YYYY-MM-DD 格式"

    p_text = str(cells[3] or "中").strip()
    priority = PRIORITY_BY_LABEL.get(p_text)
    if priority is None and p_text in ("1", "2", "3"):
        priority = int(p_text)
    if priority is None:
        return None, f"优先级「{p_text}」应为 高 / 中 / 低"

    s_text = str(cells[4] or "未完成").strip()
    status = STATUS_BY_LABEL.get(s_text)
    if status is None and s_text in ("0", "1"):
        status = int(s_text)
    if status is None:
        return None, f"完成状态「{s_text}」应为 未完成 / 已完成"

    # 课程不存在就自动建一门，不因为课程名陌生而让整行失败
    course_text = str(cells[1] or "").strip()
    course_id = (None if course_text in ("", "未分类")
                 else course_repo.get_or_create(course_text))

    item = {
        "title": title,
        "course_id": course_id,
        "due_date": due,
        "priority": priority,
        "status": status,
        "note": str(cells[5] or "").strip(),
        "completed_at": (datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                         if status == 1 else None),
    }
    return item, None


def backup_database():
    """备份数据库。

    用 SQLite 官方的 backup() API 生成一致性快照，而不是 shutil.copy 直接拷文件：
    复制正在写入的库文件有可能拿到写了一半的页，恢复时报 database disk image is malformed。
    """
    backup_dir = current_app.config["BACKUP_DIR"]
    os.makedirs(backup_dir, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    target = os.path.join(backup_dir, f"tasks_{stamp}.db")

    dest = sqlite3.connect(target)
    try:
        get_db().backup(dest)
    finally:
        dest.close()

    size_kb = round(os.path.getsize(target) / 1024, 1)
    log_repo.add_log("backup", os.path.basename(target), f"数据库快照 {size_kb} KB")
    return target


def list_backups(limit=8):
    """最近的备份文件列表，供「数据管理」页显示。"""
    backup_dir = current_app.config["BACKUP_DIR"]
    if not os.path.isdir(backup_dir):
        return []
    names = sorted((f for f in os.listdir(backup_dir) if f.endswith(".db")),
                   reverse=True)
    result = []
    for name in names[:limit]:
        path = os.path.join(backup_dir, name)
        result.append({
            "name": name,
            "size_kb": round(os.path.getsize(path) / 1024, 1),
            "mtime": datetime.fromtimestamp(
                os.path.getmtime(path)).strftime("%Y-%m-%d %H:%M:%S"),
        })
    return result
