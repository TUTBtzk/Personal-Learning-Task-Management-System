import json, sys, zipfile, re, os
from collections import Counter, defaultdict
from docx import Document
from docx.oxml.ns import qn

def fmt_len(x):
    return None if x is None else round(x.pt, 2) if hasattr(x, 'pt') else round(x.inches, 3)

def font_info(run):
    f = run.font
    rpr = run._element.rPr
    rfonts = rpr.rFonts if rpr is not None else None
    return {
        'text': run.text[:80],
        'name': f.name,
        'eastAsia': rfonts.get(qn('w:eastAsia')) if rfonts is not None else None,
        'ascii': rfonts.get(qn('w:ascii')) if rfonts is not None else None,
        'size': fmt_len(f.size),
        'bold': f.bold,
        'italic': f.italic,
        'underline': f.underline,
        'color': str(f.color.rgb) if f.color and f.color.rgb else None,
    }

def para_info(p, idx):
    pf = p.paragraph_format
    return {
        'idx': idx,
        'style': p.style.name if p.style else None,
        'text': p.text[:200].replace('\n', '\\n'),
        'align': str(p.alignment),
        'left': fmt_len(pf.left_indent),
        'right': fmt_len(pf.right_indent),
        'first': fmt_len(pf.first_line_indent),
        'before': fmt_len(pf.space_before),
        'after': fmt_len(pf.space_after),
        'line': pf.line_spacing,
        'keep_next': pf.keep_with_next,
        'keep_together': pf.keep_together,
        'runs': [font_info(r) for r in p.runs if r.text],
        'has_drawing': bool(p._element.xpath('.//w:drawing')),
    }

def style_info(s):
    pf = getattr(s, 'paragraph_format', None)
    f = s.font
    return {
        'name': s.name,
        'type': str(s.type),
        'based_on': s.base_style.name if s.base_style else None,
        'font_name': f.name,
        'font_size': fmt_len(f.size),
        'font_bold': f.bold,
        'font_italic': f.italic,
        'eastAsia': (s._element.rPr.rFonts.get(qn('w:eastAsia')) if s._element.rPr is not None and s._element.rPr.rFonts is not None else None),
        'align': str(pf.alignment) if pf else None,
        'left': fmt_len(pf.left_indent) if pf else None,
        'right': fmt_len(pf.right_indent) if pf else None,
        'first': fmt_len(pf.first_line_indent) if pf else None,
        'before': fmt_len(pf.space_before) if pf else None,
        'after': fmt_len(pf.space_after) if pf else None,
        'line': pf.line_spacing if pf else None,
    }

def summarize(path):
    doc = Document(path)
    out = {'path': os.path.abspath(path), 'sections': [], 'styles': [], 'paragraphs': [], 'tables': [], 'headers': [], 'footers': [], 'package_parts': []}
    for s in doc.sections:
        out['sections'].append({
            'start': str(s.start_type), 'orientation': str(s.orientation),
            'page_w': round(s.page_width.inches, 3), 'page_h': round(s.page_height.inches, 3),
            'left': round(s.left_margin.inches, 3), 'right': round(s.right_margin.inches, 3),
            'top': round(s.top_margin.inches, 3), 'bottom': round(s.bottom_margin.inches, 3),
            'header': round(s.header_distance.inches, 3), 'footer': round(s.footer_distance.inches, 3),
            'different_first': s.different_first_page_header_footer,
        })
    for s in doc.styles:
        if s.type in (1, 2, 3):
            out['styles'].append(style_info(s))
    out['paragraphs'] = [para_info(p, i) for i, p in enumerate(doc.paragraphs)]
    for ti, t in enumerate(doc.tables):
        out['tables'].append({'idx':ti, 'rows':len(t.rows), 'cols':len(t.columns), 'style': t.style.name if t.style else None, 'texts': [[c.text[:120] for c in row.cells] for row in t.rows[:5]]})
    for si, s in enumerate(doc.sections):
        out['headers'].append({'section':si, 'linked':s.header.is_linked_to_previous, 'paras':[para_info(p,i) for i,p in enumerate(s.header.paragraphs)]})
        out['footers'].append({'section':si, 'linked':s.footer.is_linked_to_previous, 'paras':[para_info(p,i) for i,p in enumerate(s.footer.paragraphs)]})
    with zipfile.ZipFile(path) as z:
        out['package_parts'] = sorted(z.namelist())
        out['media'] = [n for n in z.namelist() if n.startswith('word/media/')]
        out['custom_xml'] = [n for n in z.namelist() if n.startswith('customXml/')]
    return out

for path in sys.argv[1:]:
    print(json.dumps(summarize(path), ensure_ascii=False, indent=2))
    print('\n---END---\n')
